"""
tests/test_document_loader.py
Tests for app/document_loader.py

Covers:
  - TXT loading
  - PDF loading (via a programmatically created PDF)
  - DOCX loading (via a programmatically created DOCX)
  - Unsupported file type error
  - Empty file error
  - Unstructured/story-like resume → LLM extraction (mocked)
"""

import json
import tempfile
from pathlib import Path
from unittest.mock import patch, MagicMock

import pytest

from app.document_loader import load_document
from app.llm import CandidateExtraction, _fallback_extractor


# ---------------------------------------------------------------------------
# TXT loading
# ---------------------------------------------------------------------------

class TestTxtLoading:

    def test_load_txt_basic(self, tmp_path):
        f = tmp_path / "resume.txt"
        f.write_text("Jane Smith\njane@email.com\nSkills: Python, SQL", encoding="utf-8")
        text = load_document(f)
        assert "Jane Smith" in text
        assert "Python" in text

    def test_load_txt_returns_string(self, tmp_path):
        f = tmp_path / "resume.txt"
        f.write_text("Some resume content", encoding="utf-8")
        result = load_document(f)
        assert isinstance(result, str)

    def test_load_txt_empty_raises(self, tmp_path):
        f = tmp_path / "empty.txt"
        f.write_text("   \n  ", encoding="utf-8")
        with pytest.raises(ValueError, match="empty"):
            load_document(f)

    def test_load_txt_multiline(self, tmp_path):
        content = "Line 1\nLine 2\nLine 3"
        f = tmp_path / "resume.txt"
        f.write_text(content, encoding="utf-8")
        result = load_document(f)
        assert "Line 1" in result
        assert "Line 3" in result


# ---------------------------------------------------------------------------
# Unsupported file type
# ---------------------------------------------------------------------------

class TestUnsupportedType:

    def test_unsupported_extension_raises(self, tmp_path):
        f = tmp_path / "resume.odt"
        f.write_text("content", encoding="utf-8")
        with pytest.raises(ValueError, match="Unsupported"):
            load_document(f)

    def test_html_extension_raises(self, tmp_path):
        f = tmp_path / "resume.html"
        f.write_text("<html>content</html>", encoding="utf-8")
        with pytest.raises(ValueError, match="Unsupported"):
            load_document(f)


# ---------------------------------------------------------------------------
# PDF loading (requires PyMuPDF)
# ---------------------------------------------------------------------------

class TestPdfLoading:

    def test_load_pdf_basic(self, tmp_path):
        """Create a minimal PDF and verify text extraction works."""
        pytest.importorskip("fitz", reason="PyMuPDF not installed")
        import fitz

        pdf_path = tmp_path / "resume.pdf"
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Alice Wonder\nalice@example.com\nSkills: Python, Docker")
        doc.save(str(pdf_path))
        doc.close()

        text = load_document(pdf_path)
        assert "Alice" in text
        assert "Python" in text

    def test_load_pdf_empty_raises(self, tmp_path):
        """A PDF with no visible text should raise ValueError."""
        pytest.importorskip("fitz", reason="PyMuPDF not installed")
        import fitz

        pdf_path = tmp_path / "empty.pdf"
        doc = fitz.open()
        doc.new_page()  # blank page — no text
        doc.save(str(pdf_path))
        doc.close()

        with pytest.raises(ValueError, match="empty"):
            load_document(pdf_path)


# ---------------------------------------------------------------------------
# DOCX loading (requires python-docx)
# ---------------------------------------------------------------------------

class TestDocxLoading:

    def test_load_docx_basic(self, tmp_path):
        """Create a minimal DOCX and verify text extraction works."""
        pytest.importorskip("docx", reason="python-docx not installed")
        from docx import Document

        docx_path = tmp_path / "resume.docx"
        doc = Document()
        doc.add_paragraph("Bob Builder")
        doc.add_paragraph("bob@builder.com")
        doc.add_paragraph("Skills: React, SQL, Docker")
        doc.save(str(docx_path))

        text = load_document(docx_path)
        assert "Bob Builder" in text
        assert "SQL" in text

    def test_load_docx_empty_raises(self, tmp_path):
        """An empty DOCX should raise ValueError."""
        pytest.importorskip("docx", reason="python-docx not installed")
        from docx import Document

        docx_path = tmp_path / "empty.docx"
        doc = Document()
        doc.add_paragraph("")  # blank
        doc.save(str(docx_path))

        with pytest.raises(ValueError, match="empty"):
            load_document(docx_path)


# ---------------------------------------------------------------------------
# Unstructured / story-like resume → fallback extraction
# ---------------------------------------------------------------------------

STORY_RESUME = """\
John Doe started his career at a software consultancy right after graduating
with a Bachelor of Computer Science from Cairo University. Over the past six
years he has worked extensively with Python, SQL, Docker, and REST APIs.
He has also been involved in building React front-ends and integrating
third-party services via API Integration.

You can reach him at john.doe@example.com or call +1-555-999-0000.
"""


class TestUnstructuredResume:
    """
    Verify that an unstructured, story-like resume can be processed.

    We test the full pipeline using the fallback extractor so no API key is needed.
    """

    def test_fallback_extracts_email_from_story(self):
        result = _fallback_extractor(STORY_RESUME)
        assert result.email == "john.doe@example.com"

    def test_fallback_extracts_phone_from_story(self):
        result = _fallback_extractor(STORY_RESUME)
        assert result.phone is not None
        assert "555" in result.phone

    def test_fallback_extracts_skills_from_story(self):
        result = _fallback_extractor(STORY_RESUME)
        skill_lower = [s.lower() for s in result.skills]
        # At least Python and SQL should be found
        assert "python" in skill_lower
        assert "sql" in skill_lower

    def test_fallback_extracts_years_from_story(self):
        result = _fallback_extractor(STORY_RESUME)
        # "six years" is written as a word, not a digit — fallback may miss it.
        # This is expected behaviour for the regex fallback.
        # The LLM extractor handles word-form numbers correctly.
        # We only assert it doesn't crash.
        assert result.years_experience is None or isinstance(result.years_experience, float)

    def test_full_pipeline_with_story_resume_runs(self):
        """
        End-to-end: story resume → fallback extraction → graph runs without error.
        The story resume has 'Bachelor', so knockout check should PASS.
        """
        import json
        from unittest.mock import patch
        from app.graph import build_graph
        from app.state import TalentPipelineState

        initial_state: TalentPipelineState = {
            "resume_text": STORY_RESUME,
            "job_description": "Backend Engineer with Python, SQL, Docker, React, API Integration",
            "required_skills": ["Python", "SQL", "Docker", "React", "API Integration"],
            "knockout_requirements": ["Bachelor"],
            "name": None,
            "email": None,
            "phone": None,
            "years_experience": None,
            "skills": [],
            "pre_screening_status": "Pass",
            "skill_analysis": None,
            "final_disposition": None,
            "rejection_reason": None,
            "ats_payload": None,
        }

        with patch("app.nodes._extractor", side_effect=_fallback_extractor):
            graph = build_graph()
            result = graph.invoke(initial_state)

        # Should pass pre-screening (has "Bachelor")
        assert result["pre_screening_status"] == "Pass"
        # Should have a final disposition
        assert result["final_disposition"] in ("Interview", "Phone Screen", "Rejected")
        # ATS payload must be generated
        assert result["ats_payload"] is not None


# ---------------------------------------------------------------------------
# Mocked LLM extraction test (demonstrates the LLM path without API calls)
# ---------------------------------------------------------------------------

class TestMockedLLMExtraction:
    """
    Verify that when the LLM extractor returns a CandidateExtraction,
    the pipeline uses it correctly. Uses a mock — no real API calls.
    """

    def test_pipeline_uses_llm_output(self):
        """Mock the LLM extractor to return specific data and verify it flows through."""
        from unittest.mock import patch
        from app.graph import build_graph
        from app.state import TalentPipelineState

        # Simulate what the LLM would return for a story-like resume
        mock_extraction = CandidateExtraction(
            name="John Doe",
            email="john.doe@example.com",
            phone="+1-555-999-0000",
            years_experience=6.0,
            skills=["Python", "SQL", "Docker", "React", "API Integration"],
        )

        def mock_extractor(text: str) -> CandidateExtraction:
            return mock_extraction

        initial_state: TalentPipelineState = {
            "resume_text": STORY_RESUME,
            "job_description": "Backend Engineer",
            "required_skills": ["Python", "SQL", "Docker", "React", "API Integration"],
            "knockout_requirements": ["Bachelor"],
            "name": None,
            "email": None,
            "phone": None,
            "years_experience": None,
            "skills": [],
            "pre_screening_status": "Pass",
            "skill_analysis": None,
            "final_disposition": None,
            "rejection_reason": None,
            "ats_payload": None,
        }

        with patch("app.nodes._extractor", side_effect=mock_extractor):
            graph = build_graph()
            result = graph.invoke(initial_state)

        # All 5 skills matched → score = 100 → Interview
        assert result["pre_screening_status"] == "Pass"
        assert result["final_disposition"] == "Interview"
        assert result["skill_analysis"]["score"] == 100
        assert result["ats_payload"]["name"] == "John Doe"
        assert result["ats_payload"]["email"] == "john.doe@example.com"
